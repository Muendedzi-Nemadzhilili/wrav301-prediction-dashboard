import streamlit as st
import pandas as pd
import numpy as np
import matplotlib.pyplot as plt
import matplotlib.patches as mpatches
import joblib
import json
import io
from sklearn.preprocessing import StandardScaler
from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font, Alignment
from openpyxl.utils import get_column_letter

st.set_page_config(
    page_title="WRAV301 Prediction Dashboard",
    layout="wide"
)

st.markdown("""
<style>
    .main { background-color: #f8fafc; }
    .block-container { padding-top: 2rem; }
    .metric-card {
        background: white;
        border-radius: 12px;
        padding: 20px;
        text-align: center;
        box-shadow: 0 2px 8px rgba(0,0,0,0.08);
        border-left: 5px solid #1a56db;
    }
    .metric-card-danger { border-left: 5px solid #ef4444; }
    .metric-card-success { border-left: 5px solid #10b981; }
    .metric-num {
        font-size: 48px;
        font-weight: 800;
        color: #0f1e3c;
        line-height: 1;
    }
    .metric-num-danger { color: #ef4444; }
    .metric-num-success { color: #10b981; }
    .metric-label {
        font-size: 13px;
        color: #64748b;
        font-weight: 600;
        text-transform: uppercase;
        letter-spacing: 1px;
        margin-top: 6px;
    }
    .dashboard-header {
        background: linear-gradient(135deg, #0f1e3c 0%, #1a3a6e 100%);
        padding: 28px 32px;
        border-radius: 14px;
        margin-bottom: 24px;
        color: white;
    }
    .dashboard-title {
        font-size: 28px;
        font-weight: 800;
        color: white;
        margin: 0;
    }
    .dashboard-sub {
        font-size: 14px;
        color: #93c5fd;
        margin-top: 4px;
    }
    .section-header {
        font-size: 18px;
        font-weight: 700;
        color: #0f1e3c;
        margin-bottom: 12px;
        padding-bottom: 8px;
        border-bottom: 2px solid #e2e8f0;
    }
    .upload-card {
        background: white;
        border-radius: 12px;
        padding: 20px;
        box-shadow: 0 2px 8px rgba(0,0,0,0.06);
        border: 1px solid #e2e8f0;
    }
    .stButton > button {
        background: #1a56db;
        color: white;
        border-radius: 8px;
        padding: 12px 32px;
        font-weight: 700;
        font-size: 15px;
        border: none;
        width: 100%;
    }
    .warning-maroon {
        background-color: #800000;
        color: white;
        padding: 12px 16px;
        border-radius: 8px;
        font-size: 14px;
        font-weight: 500;
        margin-bottom: 8px;
    }
</style>
""", unsafe_allow_html=True)

MOODLE_FEATURES = [
    'total_events', 'unique_active_days', 'course_views',
    'module_views', 'quiz_events', 'assignment_events',
    'file_access', 'submission_events'
]
MARKS_FEATURES = ['Q_TOTAL', 'A_TOTAL', 'T1', 'T2', 'CLASS']
ALL_FEATURES   = MOODLE_FEATURES + MARKS_FEATURES

@st.cache_resource
def load_models():
    return {
        'both': {
            'model':    joblib.load('models/model_both.pkl'),
            'scaler':   joblib.load('models/scaler_both.pkl'),
            'features': json.load(open('models/features_both.json')),
        },
        'logs': {
            'model':    joblib.load('models/model_logs.pkl'),
            'scaler':   joblib.load('models/scaler_logs.pkl'),
            'features': json.load(open('models/features_logs.json')),
        },
        'marks': {
            'model':    joblib.load('models/model_marks.pkl'),
            'scaler':   joblib.load('models/scaler_marks.pkl'),
            'features': json.load(open('models/features_marks.json')),
        },
    }

models = load_models()

def extract_log_features(logs_df, student_ids):
    logs_df['Time'] = pd.to_datetime(logs_df['Time'], format='mixed', dayfirst=False, errors='coerce')
    logs_df = logs_df.dropna(subset=['Time'])
    logs_df['Date'] = logs_df['Time'].dt.date
    features = logs_df.groupby('User full name').agg(
        total_events       = ('Event name', 'count'),
        unique_active_days = ('Date', 'nunique'),
        course_views       = ('Event name', lambda x: (x == 'Course viewed').sum()),
        module_views       = ('Event name', lambda x: (x == 'Course module viewed').sum()),
        quiz_events        = ('Component',  lambda x: (x == 'Quiz').sum()),
        assignment_events  = ('Component',  lambda x: (x == 'Assignment').sum()),
        file_access        = ('Component',  lambda x: (x == 'File').sum()),
        submission_events  = ('Event name', lambda x: x.str.contains('ubmission', case=False).sum()),
    ).reset_index().rename(columns={'User full name': 'Student ID'})
    base   = pd.DataFrame({'Student ID': student_ids})
    merged = base.merge(features, on='Student ID', how='left')
    merged[MOODLE_FEATURES] = merged[MOODLE_FEATURES].fillna(0)
    return merged

def build_feature_matrix(logs_df, marks_df):
    missing_mark_cols = []
    missing_log = False

    if marks_df is not None and 'Student ID' in marks_df.columns:
        student_ids = marks_df['Student ID'].tolist()
    elif logs_df is not None:
        student_ids = logs_df['User full name'].unique().tolist()
    else:
        return None, [], []

    result = pd.DataFrame({'Student ID': student_ids})

    if logs_df is not None:
        log_feats = extract_log_features(logs_df, student_ids)
        result = result.merge(log_feats, on='Student ID', how='left')
        result[MOODLE_FEATURES] = result[MOODLE_FEATURES].fillna(0)
    else:
        for col in MOODLE_FEATURES:
            result[col] = 0
        missing_log = True

    if marks_df is not None:
        available_marks = [c for c in MARKS_FEATURES if c in marks_df.columns]
        missing_mark_cols = [c for c in MARKS_FEATURES if c not in marks_df.columns]
        marks_cols = ['Student ID'] + available_marks
        if 'FINAL' in marks_df.columns:
            marks_cols.append('FINAL')
        result = result.merge(marks_df[marks_cols], on='Student ID', how='left')
        for col in MARKS_FEATURES:
            if col not in result.columns:
                result[col] = 0
            else:
                result[col] = result[col].fillna(0)
    else:
        for col in MARKS_FEATURES:
            result[col] = 0

    return result, missing_mark_cols, missing_log

def select_model(logs_df, marks_df, missing_mark_cols):
    has_logs  = logs_df is not None
    has_marks = marks_df is not None
    all_marks_present = len(missing_mark_cols) == 0

    if has_logs and has_marks and all_marks_present:
        return 'both'
    elif has_logs and not has_marks:
        return 'logs'
    elif has_marks and not has_logs:
        return 'marks'
    else:
        return 'both'

def generate_excel(results_df):
    wb = Workbook()
    ws = wb.active
    ws.title = "Prediction Results"
    red         = PatternFill(start_color="FFCCCC", end_color="FFCCCC", fill_type="solid")
    green       = PatternFill(start_color="CCFFCC", end_color="CCFFCC", fill_type="solid")
    header_fill = PatternFill(start_color="0F1E3C", end_color="0F1E3C", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True)
    ws.append(list(results_df.columns))
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")
    for _, row in results_df.iterrows():
        ws.append(list(row))
        fill = red if row['Prediction'] == 'At-Risk' else green
        for cell in ws[ws.max_row]:
            cell.fill = fill
            cell.alignment = Alignment(horizontal="center")
    for col in ws.columns:
        max_len = max(len(str(cell.value or "")) for cell in col)
        ws.column_dimensions[get_column_letter(col[0].column)].width = max_len + 4
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output

st.markdown("""
<div class="dashboard-header">
    <div class="dashboard-title">WRAV301 Prediction Dashboard</div>
    <div class="dashboard-sub">Nelson Mandela University &nbsp;·&nbsp; Advanced Data Structures</div>
</div>
""", unsafe_allow_html=True)

st.markdown('<div class="section-header">Upload Data Files</div>', unsafe_allow_html=True)
st.info("You can upload either or both files but prediction accuracy is best when all data is available.")

col1, col2 = st.columns(2)
with col1:
    st.markdown('<div class="upload-card">', unsafe_allow_html=True)
    st.markdown("**Moodle Log File** *(optional)*")
    logs_file = st.file_uploader("Upload logs CSV", type="csv", key="logs")
    if logs_file:
        st.success(f"{logs_file.name} uploaded")
    st.markdown('</div>', unsafe_allow_html=True)

with col2:
    st.markdown('<div class="upload-card">', unsafe_allow_html=True)
    st.markdown("**Student Marks File** *(optional)*")
    marks_file = st.file_uploader("Upload marks CSV", type="csv", key="marks")
    if marks_file:
        st.success(f"{marks_file.name} uploaded")
    st.markdown('</div>', unsafe_allow_html=True)

if not logs_file and not marks_file:
    st.info("Please upload at least one file above to get started.")
    st.stop()

logs_df  = pd.read_csv(logs_file)  if logs_file  else None
marks_df = pd.read_csv(marks_file) if marks_file else None

if logs_df is not None:
    with st.expander("Preview Moodle Logs"):
        st.dataframe(logs_df.head(), use_container_width=True)

if marks_df is not None:
    with st.expander("Preview Student Marks"):
        st.dataframe(marks_df.head(), use_container_width=True)

missing_mark_cols = [c for c in MARKS_FEATURES if marks_df is not None and c not in marks_df.columns]
missing_log = logs_df is None

warnings_to_show = []

if missing_mark_cols:
    available = [c for c in MARKS_FEATURES if marks_df is not None and c in marks_df.columns]
    warnings_to_show.append(
        f"The following assessment columns are missing from the marks file: "
        f"{', '.join(missing_mark_cols)}. These will be set to zero. "
        f"Predictions may be less accurate.<br>"
        f"Assessment columns found: {', '.join(available)}"
    )

if missing_log:
    warnings_to_show.append(
        "No Moodle log file uploaded. All Moodle engagement features will be "
        "set to zero. Prediction accuracy will be reduced."
    )

for w in warnings_to_show:
    st.markdown(
        f'<div class="warning-maroon">{w}</div>',
        unsafe_allow_html=True
    )

st.markdown("<br>", unsafe_allow_html=True)
run = st.button("Run Prediction", type="primary")

if run:
    with st.spinner("Extracting features and running prediction model..."):
        merged, missing_mark_cols, missing_log = build_feature_matrix(logs_df, marks_df)

        if merged is None:
            st.error("Could not build feature matrix. Please upload at least one valid file.")
            st.stop()

        mode     = select_model(logs_df, marks_df, missing_mark_cols)
        selected = models[mode]
        feature_cols = selected['features']

        for col in feature_cols:
            if col not in merged.columns:
                merged[col] = 0

        X        = merged[feature_cols]
        X_scaled = selected['scaler'].transform(X)
        preds    = selected['model'].predict(X_scaled)
        merged['Prediction'] = ['At-Risk' if p == 1 else 'Not At-Risk' for p in preds]

    total    = len(merged)
    at_risk  = int(sum(preds == 1))
    not_risk = int(sum(preds == 0))
    pct      = round(at_risk / total * 100, 1)

    st.markdown("---")
    st.markdown('<div class="section-header">Prediction Results</div>', unsafe_allow_html=True)

    c1, c2, c3, c4 = st.columns(4)
    with c1:
        st.markdown(f"""
        <div class="metric-card">
            <div class="metric-num">{total}</div>
            <div class="metric-label">Total Students</div>
        </div>""", unsafe_allow_html=True)
    with c2:
        st.markdown(f"""
        <div class="metric-card metric-card-danger">
            <div class="metric-num metric-num-danger">{at_risk}</div>
            <div class="metric-label">At-Risk</div>
        </div>""", unsafe_allow_html=True)
    with c3:
        st.markdown(f"""
        <div class="metric-card metric-card-success">
            <div class="metric-num metric-num-success">{not_risk}</div>
            <div class="metric-label">Not At-Risk</div>
        </div>""", unsafe_allow_html=True)
    with c4:
        st.markdown(f"""
        <div class="metric-card">
            <div class="metric-num" style="color:#1a56db">{pct}%</div>
            <div class="metric-label">At-Risk Rate</div>
        </div>""", unsafe_allow_html=True)

    st.markdown("<br>", unsafe_allow_html=True)

    ch1, ch2 = st.columns(2)

    with ch1:
        st.markdown("**Risk Distribution**")
        fig1, ax1 = plt.subplots(figsize=(5, 4))
        bars = ax1.bar(
            ['Not At-Risk', 'At-Risk'], [not_risk, at_risk],
            color=['#10b981', '#ef4444'], width=0.5,
            edgecolor='white', linewidth=1.5
        )
        ax1.set_ylabel('Number of Students', fontsize=11)
        ax1.set_title('Prediction Summary', fontsize=12, fontweight='bold', pad=12)
        ax1.spines[['top', 'right']].set_visible(False)
        ax1.set_facecolor('#f8fafc')
        fig1.patch.set_facecolor('#f8fafc')
        for bar in bars:
            h = bar.get_height()
            ax1.text(
                bar.get_x() + bar.get_width() / 2, h + 0.3,
                str(int(h)), ha='center', fontweight='bold', fontsize=12
            )
        plt.tight_layout()
        st.pyplot(fig1)

    with ch2:
        st.markdown("**Feature Importance**")
        importance = pd.Series(
            selected['model'].feature_importances_,
            index=feature_cols
        ).sort_values(ascending=True)
        fig2, ax2 = plt.subplots(figsize=(5, 4))
        colors = [
            '#ef4444' if i >= len(importance) - 3 else '#1a56db'
            for i in range(len(importance))
        ]
        importance.plot(kind='barh', ax=ax2, color=colors, edgecolor='white')
        ax2.set_title('Feature Importance — Random Forest',
                      fontsize=12, fontweight='bold', pad=12)
        ax2.set_xlabel('Importance Score', fontsize=10)
        ax2.spines[['top', 'right']].set_visible(False)
        ax2.set_facecolor('#f8fafc')
        fig2.patch.set_facecolor('#f8fafc')
        ax2.legend(handles=[
            mpatches.Patch(color='#ef4444', label='Top 3 features'),
            mpatches.Patch(color='#1a56db', label='Other features')
        ], fontsize=9)
        plt.tight_layout()
        st.pyplot(fig2)

    st.markdown("<br>", unsafe_allow_html=True)
    st.markdown("**At-Risk Students**")

    marks_show  = ['FINAL'] if marks_df is not None and 'FINAL' in merged.columns else []
    moodle_show = ['total_events', 'unique_active_days', 'quiz_events', 'assignment_events'] if logs_df is not None else []
    display_cols = ['Student ID'] + marks_show + moodle_show + ['Prediction']
    display_cols = [c for c in display_cols if c in merged.columns]

    at_risk_df = merged[merged['Prediction'] == 'At-Risk'][display_cols].reset_index(drop=True)

    st.dataframe(
        at_risk_df.style.applymap(
            lambda v: 'background-color: #ffcccc; color: #991b1b; font-weight:bold'
            if v == 'At-Risk' else '',
            subset=['Prediction']
        ),
        use_container_width=True
    )

    st.markdown("<br>", unsafe_allow_html=True)
    st.markdown("---")
    st.markdown('<div class="section-header">Download Report</div>', unsafe_allow_html=True)

    excel_data = generate_excel(merged[display_cols])

    dl_col, show_col = st.columns([1, 1])
    with dl_col:
        st.download_button(
            label="Download Full Excel Report",
            data=excel_data,
            file_name="WRAV301_Prediction_Report.xlsx",
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
    with show_col:
        show_table = st.button("Show Full Results Table")

    if show_table:
        st.markdown("<br>", unsafe_allow_html=True)
        st.markdown("**Full Student Results**")
        st.dataframe(
            merged[display_cols].style.applymap(
                lambda v: 'background-color: #ffcccc; color: #991b1b; font-weight:bold'
                if v == 'At-Risk'
                else ('background-color: #ccffcc; color: #166534; font-weight:bold'
                      if v == 'Not At-Risk' else ''),
                subset=['Prediction']
            ),
            use_container_width=True,
            height=600
        )